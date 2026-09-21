"""
Gas update PREFLIGHT - validate the gas master's STRUCTURE before extracting.

Run before every refresh. Locates the bcf block (MASTER tab) and the generation
block (Ember Electricity tab) by anchor, and compares the row labels + date span
against the last known-good manifest (WORKING/gas_manifest.json):

  exit 0  -> SAFE to proceed. No structural change, or only NEW rows (they flow
             into the tables automatically; the names are printed because the
             chart selector groups in generate_dashboard.py list rows explicitly,
             so a new import source / consumption item won't be selectable in
             the range chart until Claude adds it).
  exit 2  -> STOP, needs review: a missing anchor / layout change, or a REMOVED
             or RENAMED row (the dashboard's chart groups and colours are keyed
             by exact label). Call Claude before updating.

Also prints an informational "what changed" summary against the previous
extract (WORKING/gas_model_input.xlsx): the first month whose values differ per
block and the new last month of data - the "did I paste the right file?" check.

  py -3 src/preflight_gas.py            # check against the manifest
  py -3 src/preflight_gas.py --write    # (re)initialise the manifest from the master
                                        #   (done automatically after an approved update)
"""
import json
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

from extract_gas_input import (SRC, OUT, MASTER_SHEET, POWER_SHEET, LayoutError,
                               _clean, _read_rows, locate_master, locate_power, open_master)

ROOT = Path(__file__).resolve().parent.parent
MANIFEST = ROOT / "WORKING" / "gas_manifest.json"


def snapshot():
    """Structural fingerprint of the master: row labels of both blocks + date span."""
    wb = open_master()
    try:
        for tab in (MASTER_SHEET, POWER_SHEET):
            if tab not in wb.sheetnames:
                raise LayoutError(f"tab '{tab}' not found in the master")
        ws = wb[MASTER_SHEET]
        m = locate_master(ws)
        labels = _read_rows(ws, m["first_row"], m["last_row"], 1)
        gas_rows = [_clean(v[0]) for v in labels.values()]
        dates = next(ws.iter_rows(min_row=m["date_row"], max_row=m["date_row"],
                                  max_col=m["last_col"], values_only=True))

        wp = wb[POWER_SHEET]
        p = locate_power(wp)
        plabels = _read_rows(wp, p["first_row"], p["last_row"], 1)
        power_rows = [_clean(v[0]) for v in plabels.values()]
        pdates = next(wp.iter_rows(min_row=p["date_row"], max_row=p["date_row"],
                                   max_col=p["last_col"], values_only=True))
    finally:
        wb.close()
    return {
        "written": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "gas": {"rows": gas_rows, "first_month": dates[1].strftime("%Y-%m"),
                "last_month": dates[-1].strftime("%Y-%m"), "block": [m["first_row"], m["last_row"]]},
        "power": {"rows": power_rows, "first_month": pdates[1].strftime("%Y-%m"),
                  "last_month": pdates[-1].strftime("%Y-%m"), "block": [p["first_row"], p["last_row"]]},
    }


def diff_rows(name, old, new):
    """Compare two ordered label lists. Returns (stop, notes)."""
    stop, notes = False, []
    so, sn = set(old), set(new)
    for r in new:
        if r not in so:
            notes.append(f"  + NEW ROW    ({name}): '{r}'")
    for r in old:
        if r not in sn:
            notes.append(f"  - REMOVED/RENAMED ({name}): '{r}'")
            stop = True
    if not stop and old != new and so == sn:
        notes.append(f"  ~ ROW ORDER changed ({name}) - tables follow the master's order")
    return stop, notes


def value_change_report(snap):
    """Where do the master's numbers differ from the previous extract? (informational)"""
    if not OUT.exists():
        print("  (no previous extract to compare against - first run)")
        return
    try:
        prev = openpyxl.load_workbook(OUT, read_only=True, data_only=True)
        wb = open_master()
    except Exception as e:                       # never block on this
        print(f"  (could not compare with previous extract: {e})")
        return
    try:
        blocks = [
            ("Gas balance", wb[MASTER_SHEET], locate_master(wb[MASTER_SHEET]), prev["Monthly Data"], 1, 5),
            ("Power",       wb[POWER_SHEET],  locate_power(wb[POWER_SHEET]),  prev["Power Data"],   6, 8),
        ]
        for name, ws, loc, wprev, prev_date_row, prev_first in blocks:
            cur = _read_rows(ws, 1, loc["last_row"], loc["last_col"])
            old = {i: list(r) for i, r in enumerate(wprev.iter_rows(values_only=True), 1)}
            dates = cur[loc["date_row"]]
            old_by_label = {}
            for i in range(prev_first, len(old) + 1):
                lab = _clean(old[i][0]) if old[i] else ""
                if lab:
                    old_by_label[lab] = old[i]
            old_last = old.get(prev_date_row, [None])
            old_last_d = next((v for v in reversed(old_last) if isinstance(v, datetime)), None)
            first_diff, n_rows = None, 0
            for r in range(loc["first_row"], loc["last_row"] + 1):
                lab = _clean(cur[r][0])
                o = old_by_label.get(lab)
                if o is None:
                    continue
                for c in range(1, min(len(cur[r]), len(o))):
                    a, b = cur[r][c], o[c]
                    if (a or 0) != (b or 0) and abs((a or 0) - (b or 0)) > 1e-9:
                        d = dates[c]
                        if first_diff is None or d < first_diff:
                            first_diff = d
                        n_rows += 1
                        break
            span = f"{dates[1]:%b %Y} - {dates[-1]:%b %Y}"
            if first_diff is None:
                print(f"  {name}: NO value changes vs the previous extract ({span})")
            else:
                print(f"  {name}: {n_rows} row(s) changed, earliest change {first_diff:%b %Y}; "
                      f"data now {span}")
    except Exception as e:
        print(f"  (comparison skipped: {e})")
    finally:
        prev.close()
        wb.close()


def main():
    write = "--write" in sys.argv
    print("=" * 60)
    print("GAS PREFLIGHT" + ("  (writing manifest)" if write else ""))
    print("=" * 60)
    if not SRC.exists():
        print(f"*** Master not found: {SRC}")
        print("*** Put the latest 'AKAP Global Gas Model.xlsx' in the INPUT folder.")
        return 2
    try:
        snap = snapshot()
    except LayoutError as e:
        print(f"*** LAYOUT ERROR: {e}")
        print("*** The master's layout has changed - call Claude before updating.")
        return 2

    g, p = snap["gas"], snap["power"]
    print(f"Master OK: gas balance {len(g['rows'])} rows ({g['first_month']} .. {g['last_month']}), "
          f"power {len(p['rows'])} rows ({p['first_month']} .. {p['last_month']})")

    if write:
        MANIFEST.parent.mkdir(parents=True, exist_ok=True)
        MANIFEST.write_text(json.dumps(snap, indent=2), encoding="utf-8")
        print(f"Manifest written: {MANIFEST}")
        return 0

    if not MANIFEST.exists():
        print("No manifest yet (first run) - structure recorded after this build.")
        print("\nChanges vs the previous extract:")
        value_change_report(snap)
        return 0

    old = json.loads(MANIFEST.read_text(encoding="utf-8"))
    stop, notes = False, []
    for key, name in (("gas", "gas balance"), ("power", "power")):
        s, n = diff_rows(name, old[key]["rows"], snap[key]["rows"])
        stop, notes = stop or s, notes + n
        if old[key]["last_month"] != snap[key]["last_month"]:
            notes.append(f"  ~ {name}: last month {old[key]['last_month']} -> {snap[key]['last_month']}")
        if old[key]["first_month"] != snap[key]["first_month"]:
            notes.append(f"  ~ {name}: first month {old[key]['first_month']} -> {snap[key]['first_month']}")

    print(f"\nStructure vs last approved update ({old.get('written', '?')}):")
    if notes:
        print("\n".join(notes))
    else:
        print("  no structural changes")
    if any(n.lstrip().startswith("+ NEW ROW") for n in notes):
        print("  NOTE: new rows appear in the tables automatically, but the range-chart")
        print("        selector lists rows by name - ask Claude to add them to a group.")

    print("\nChanges vs the previous extract:")
    value_change_report(snap)

    if stop:
        print("\nRESULT: STOP - a row was removed or renamed. Nothing has been changed.")
        return 2
    print("\nRESULT: PROCEED")
    return 0


if __name__ == "__main__":
    sys.exit(main())
