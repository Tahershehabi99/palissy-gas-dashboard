"""
Extract the European gas balance + power generation data from the AKAP Global
Gas Model master into the clean input workbook the dashboard reads.

Source : INPUT/AKAP Global Gas Model.xlsx  (the user drops the latest master here;
         READ-ONLY - opened but never modified/saved)
Output : WORKING/gas_model_input.xlsx      (auto-generated clean extract)

Everything is located by ANCHOR (row labels in column A), never by fixed row
numbers, so a row inserted in the master does not silently shift the data:

  'MASTER' tab  -> 'Monthly Data' sheet
      'Days' row + the month-date row (first row with a real date in col B)
      the 'bcf' block: from the row after the 'bcf' header down to
      'Storage percentage' (inclusive). 'Stock change' and anything below are
      not copied (the dashboard derives it).
      Output layout: row 1 Date, row 2 Days, row 3 blank, row 4 'bcf', data from row 5.

  'Ember Electricity' tab -> 'Power Data' sheet
      the header rows (everything above the block header) copied as-is, then
      'Generation by Source - EU + UK' and its rows down to the first blank label.
      Output layout: identical row positions to the master (date row 6, days row 4,
      block header row 7, data from row 8) - generate_dashboard.py reads it that way.

  cell COLOURS -> 'Boundaries' sheet   (Sheet | Row | Last actual month)
      The master marks actual vs forecast by colour: black = reported, blue or
      red-on-yellow = forecast. Per row, the last actual month is the month before
      the first coloured RUN (>= MIN_RUN consecutive coloured cells - a lone
      coloured cell is a note, and the Ember tab colours only the current forecast
      year). Rows that are entirely black (derived rows) get no entry. The
      dashboard shades each row from its own boundary; set_gas_actual.py uses the
      same data to propose the tab-level boundary.

Usage:  py -3 src/extract_gas_input.py
"""
import os
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Border, Side

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "INPUT" / "AKAP Global Gas Model.xlsx"
OUT = ROOT / "WORKING" / "gas_model_input.xlsx"

MASTER_SHEET = "MASTER"
POWER_SHEET = "Ember Electricity"
POWER_HEADER = "Generation by Source - EU + UK"
BCF_END_LABEL = "Storage percentage"     # last row of the bcf block (inclusive)
BCF_STOP_LABEL = "Stock change"          # never copied; marks the end if seen first
SCAN_ROWS = 80                           # how far down col A we look for anchors
MIN_RUN = 3                              # coloured cells needed to count as "forecast starts here"


class LayoutError(Exception):
    """An anchor we rely on is missing - the master's layout changed."""


def _clean(v):
    return str(v).strip() if v is not None else ""


def _read_rows(ws, r0, r1, max_col):
    """rows r0..r1 (inclusive) as lists of values, via iter_rows (fast in read_only)."""
    out = {}
    for i, row in enumerate(ws.iter_rows(min_row=r0, max_row=r1, max_col=max_col, values_only=True), r0):
        out[i] = list(row)
    return out


def _last_date_col(row_vals):
    """1-based index of the last real date in a header row."""
    last = None
    for i, v in enumerate(row_vals, 1):
        if isinstance(v, datetime):
            last = i
    return last


def locate_master(ws):
    """Find the Days row, the date row and the bcf block on the MASTER tab.
    Returns dict(days_row, date_row, bcf_header_row, first_row, last_row, last_col)."""
    head = _read_rows(ws, 1, SCAN_ROWS, 2)          # col A labels + col B sample
    days_row = date_row = bcf_row = None
    for r, (a, b) in head.items():
        la = _clean(a).lower()
        if days_row is None and la == "days":
            days_row = r
        if date_row is None and isinstance(b, datetime):
            date_row = r
        if bcf_row is None and la == "bcf":
            bcf_row = r
    if days_row is None:
        raise LayoutError("MASTER: no 'Days' row found in column A")
    if date_row is None:
        raise LayoutError("MASTER: no month-date row found (no date in column B)")
    if bcf_row is None:
        raise LayoutError("MASTER: no 'bcf' section header found in column A")

    first_row = bcf_row + 1
    last_row = None
    for r in range(first_row, SCAN_ROWS + 1):
        la = _clean(head.get(r, [None, None])[0])
        if la == BCF_STOP_LABEL or la == "":
            break
        if la == BCF_END_LABEL:
            last_row = r
            break
    if last_row is None:
        raise LayoutError(f"MASTER: bcf block does not end with '{BCF_END_LABEL}' "
                          f"(looked from row {first_row})")

    date_vals = next(ws.iter_rows(min_row=date_row, max_row=date_row, max_col=600, values_only=True))
    last_col = _last_date_col(date_vals)
    if last_col is None or last_col < 2:
        raise LayoutError("MASTER: date row has no dates")
    return {"days_row": days_row, "date_row": date_row, "bcf_header_row": bcf_row,
            "first_row": first_row, "last_row": last_row, "last_col": last_col}


def locate_power(ws):
    """Find the generation block on the Ember tab.
    Returns dict(days_row, date_row, header_row, first_row, last_row, last_col)."""
    head = _read_rows(ws, 1, SCAN_ROWS, 2)
    days_row = date_row = header_row = None
    for r, (a, b) in head.items():
        la = _clean(a).lower()
        if days_row is None and la == "days":
            days_row = r
        if date_row is None and isinstance(b, datetime):
            date_row = r
        if header_row is None and _clean(a) == POWER_HEADER:
            header_row = r
    if days_row is None:
        raise LayoutError(f"{POWER_SHEET}: no 'Days' row found in column A")
    if date_row is None:
        raise LayoutError(f"{POWER_SHEET}: no month-date row found (no date in column B)")
    if header_row is None:
        raise LayoutError(f"{POWER_SHEET}: no '{POWER_HEADER}' header found in column A")
    if not (days_row < header_row and date_row < header_row):
        raise LayoutError(f"{POWER_SHEET}: Days/date rows are not above the generation block")

    first_row = header_row + 1
    last_row = None
    for r in range(first_row, SCAN_ROWS + 1):
        if _clean(head.get(r, [None, None])[0]) == "":
            last_row = r - 1
            break
    if last_row is None or last_row < first_row:
        raise LayoutError(f"{POWER_SHEET}: generation block is empty")

    date_vals = next(ws.iter_rows(min_row=date_row, max_row=date_row, max_col=600, values_only=True))
    last_col = _last_date_col(date_vals)
    if last_col is None or last_col < 2:
        raise LayoutError(f"{POWER_SHEET}: date row has no dates")
    return {"days_row": days_row, "date_row": date_row, "header_row": header_row,
            "first_row": first_row, "last_row": last_row, "last_col": last_col}


def _is_black(cell):
    """True when the cell is plain (no font colour / theme text colour, no fill)."""
    f = cell.font
    col = f.color if f is not None else None
    if col is not None:
        if col.type == "rgb" and isinstance(col.rgb, str) and col.rgb not in ("FF000000", "00000000"):
            return False
        if col.type == "theme" and col.theme not in (0, 1):      # 0/1 = window/text
            return False
    fl = cell.fill
    if fl is not None and fl.fill_type == "solid":
        fg = fl.fgColor.rgb if fl.fgColor is not None else None
        if isinstance(fg, str) and fg not in ("00000000", "FFFFFFFF"):
            return False
    return True


def _row_boundary(cells, dates):
    """Last actual month (datetime) for one row of cells, or None when the row is
    all black (no information) or coloured from its first value."""
    flags = []
    for i, c in enumerate(cells):
        if i == 0 or not isinstance(dates[i], datetime) or c.value is None:
            flags.append(None)
        else:
            flags.append(_is_black(c))
    n = len(flags)
    for i in range(1, n):
        if flags[i] is False:
            run = 0
            for j in range(i, n):
                if flags[j] is False:
                    run += 1
                elif flags[j] is True:
                    break
            if run >= MIN_RUN:
                for k in range(i - 1, 0, -1):          # previous month WITH a value
                    if flags[k] is True:
                        return dates[k]
                return None
    return None


def row_boundaries(wb):
    """{'gas': {label: 'YYYY-MM'}, 'power': {...}} read from the master's cell colours.
    Only rows with a detectable boundary are present."""
    out = {}
    for key, sheet, locate in (("gas", MASTER_SHEET, locate_master), ("power", POWER_SHEET, locate_power)):
        ws = wb[sheet]
        loc = locate(ws)
        rows = list(ws.iter_rows(min_row=1, max_row=loc["last_row"], max_col=loc["last_col"]))
        dates = [c.value for c in rows[loc["date_row"] - 1]]
        found = {}
        for r in range(loc["first_row"], loc["last_row"] + 1):
            cells = rows[r - 1]
            d = _row_boundary(cells, dates)
            if d is not None:
                found[_clean(cells[0].value)] = d.strftime("%Y-%m")
        out[key] = found
    return out


def open_master():
    if not SRC.exists():
        raise FileNotFoundError(f"Master not found: {SRC}\n  Put the latest 'AKAP Global Gas Model.xlsx' in the INPUT folder.")
    # read_only + iter_rows = fast; the file is 8 MB of formulas otherwise.
    return openpyxl.load_workbook(SRC, read_only=True, data_only=True)


def main():
    print(f"Reading master: {SRC.name}")
    wb = open_master()
    if MASTER_SHEET not in wb.sheetnames:
        raise LayoutError(f"tab '{MASTER_SHEET}' not found in the master")
    if POWER_SHEET not in wb.sheetnames:
        raise LayoutError(f"tab '{POWER_SHEET}' not found in the master")

    # ---------------- gas balance (MASTER -> Monthly Data) ----------------
    ws = wb[MASTER_SHEET]
    m = locate_master(ws)
    rows = _read_rows(ws, 1, m["last_row"], m["last_col"])
    n_months = m["last_col"] - 1
    first_d, last_d = rows[m["date_row"]][1], rows[m["date_row"]][m["last_col"] - 1]
    print(f"  MASTER: bcf block rows {m['first_row']}-{m['last_row']} "
          f"({m['last_row'] - m['first_row'] + 1} rows), {n_months} months "
          f"({first_d:%b %Y} - {last_d:%b %Y})")

    out = openpyxl.Workbook()
    ws_out = out.active
    ws_out.title = "Monthly Data"

    header_font = Font(name="Calibri", size=10, bold=True)
    data_font = Font(name="Calibri", size=10)
    date_font = Font(name="Calibri", size=9, bold=True)
    label_font = Font(name="Calibri", size=10, bold=False)
    parent_font = Font(name="Calibri", size=10, bold=True)
    thin_border = Border(bottom=Side(style="thin", color="D0D0D0"))

    ws_out.cell(row=1, column=1, value="Date").font = header_font
    ws_out.cell(row=2, column=1, value="Days").font = header_font
    ws_out.cell(row=4, column=1, value="bcf").font = Font(name="Calibri", size=10, bold=True, italic=True)
    for c in range(2, m["last_col"] + 1):
        d = rows[m["date_row"]][c - 1]
        cell = ws_out.cell(row=1, column=c, value=d)
        cell.font = date_font
        if d is not None:
            cell.number_format = "MMM-YY"
        cell = ws_out.cell(row=2, column=c, value=rows[m["days_row"]][c - 1])
        cell.font = data_font

    stock_labels = {"Opening Storage", "Closing Storage", "Storage percentage"}
    r_out = 5
    for r in range(m["first_row"], m["last_row"] + 1):
        src = rows[r]
        label = _clean(src[0])
        is_parent = label.startswith("+") or label.startswith("-") or label in stock_labels
        cell = ws_out.cell(row=r_out, column=1, value=label)
        cell.font = parent_font if is_parent else label_font
        cell.border = thin_border
        for c in range(2, m["last_col"] + 1):
            v = src[c - 1]
            cell = ws_out.cell(row=r_out, column=c, value=v)
            cell.font = data_font
            cell.border = thin_border
            if v is not None:
                cell.number_format = "0.0%" if label == BCF_END_LABEL else "#,##0.0"
        r_out += 1
    ws_out.column_dimensions["A"].width = 30
    for c in range(2, m["last_col"] + 1):
        ws_out.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 12
    ws_out.freeze_panes = "B2"
    gas_rows = r_out - 5

    # ---------------- instructions sheet ----------------
    ws_inst = out.create_sheet("Instructions")
    for i, line in enumerate([
        "PALISSY GAS MODEL - INPUT FILE  (auto-generated - do not edit)",
        "",
        "This file is rebuilt from INPUT/AKAP Global Gas Model.xlsx every time you run",
        "INPUT/update_gas.bat. Any manual edits here are overwritten.",
        "",
        "TO UPDATE THE DASHBOARD:",
        "1. Put the latest 'AKAP Global Gas Model.xlsx' in the INPUT folder",
        "2. Double-click INPUT/update_gas.bat and follow the prompts",
        "3. Review output/index.html, then run the push bat in the output folder",
        "",
        "'Monthly Data' = the bcf block of the MASTER tab (Opening Storage .. Storage percentage)",
        "'Power Data'   = the 'Generation by Source - EU + UK' block of the Ember Electricity tab",
    ], 1):
        ws_inst.cell(row=i, column=1, value=line).font = Font(name="Calibri", size=11)
    ws_inst.column_dimensions["A"].width = 90

    # ---------------- power (Ember Electricity -> Power Data) ----------------
    wp = wb[POWER_SHEET]
    p = locate_power(wp)
    prow = _read_rows(wp, 1, p["last_row"], p["last_col"])
    pn = p["last_col"] - 1
    pf, pl = prow[p["date_row"]][1], prow[p["date_row"]][p["last_col"] - 1]
    print(f"  {POWER_SHEET}: block rows {p['first_row']}-{p['last_row']} "
          f"({p['last_row'] - p['first_row'] + 1} rows), {pn} months ({pf:%b %Y} - {pl:%b %Y})")

    ws_p = out.create_sheet("Power Data")
    for r in range(1, p["last_row"] + 1):
        for c in range(1, p["last_col"] + 1):
            v = prow[r][c - 1]
            if v is not None:
                cell = ws_p.cell(row=r, column=c, value=v)
                if isinstance(v, datetime):
                    cell.number_format = "MMM-YY"
    ws_p.column_dimensions["A"].width = 38
    ws_p.freeze_panes = "B7"

    # ---------------- per-row actual/forecast boundaries (from cell colours) ----------------
    bounds = row_boundaries(wb)
    ws_b = out.create_sheet("Boundaries")
    for c, h in enumerate(("Sheet", "Row", "Last actual month"), 1):
        ws_b.cell(row=1, column=c, value=h).font = header_font
    rb = 2
    for key, sheet_name in (("gas", "Monthly Data"), ("power", "Power Data")):
        for label, ym in bounds[key].items():
            ws_b.cell(row=rb, column=1, value=sheet_name)
            ws_b.cell(row=rb, column=2, value=label)
            ws_b.cell(row=rb, column=3, value=ym)
            rb += 1
    ws_b.column_dimensions["A"].width = 14
    ws_b.column_dimensions["B"].width = 38
    ws_b.column_dimensions["C"].width = 18
    n_b = {k: len(v) for k, v in bounds.items()}

    wb.close()
    os.makedirs(OUT.parent, exist_ok=True)
    out.save(OUT)
    print(f"Saved: {OUT}")
    print(f"  Monthly Data: {gas_rows} rows x {n_months} months")
    print(f"  Power Data  : {p['last_row'] - p['first_row'] + 1} rows x {pn} months")
    print(f"  Boundaries  : colour-detected last actual month for {n_b['gas']} gas + {n_b['power']} power rows")


if __name__ == "__main__":
    try:
        main()
    except (LayoutError, FileNotFoundError) as e:
        print(f"\n*** ERROR: {e}")
        raise SystemExit(1)
